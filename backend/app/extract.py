from pathlib import Path


def extract_document(filename: str, content: bytes) -> tuple[str, list[dict]]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        text = content.decode("utf-8-sig", errors="replace")
        return "document", [{"page": 1, "text": text}]
    if suffix == ".pdf":
        from pypdf import PdfReader
        import io
        pages = [{"page": i + 1, "text": page.extract_text() or ""} for i, page in enumerate(PdfReader(io.BytesIO(content)).pages)]
        return "pdf", pages
    if suffix == ".docx":
        from docx import Document
        import io
        doc = Document(io.BytesIO(content))
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append((para.style.name if para.style else "", text))
        body = "\n\n".join((f"{text}:" if "Heading" in style else text) for style, text in paragraphs)
        for table in doc.tables:
            body += "\n\n" + "\n".join(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return "docx", [{"page": 1, "text": body}]
    raise ValueError("Unsupported file type. Upload PDF, DOCX, TXT, or Markdown.")
